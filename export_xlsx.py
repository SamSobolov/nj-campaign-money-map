"""Export the site data to a workbook with one tab per committee per quarter (for the Google Sheet).

Usage: python3 export_xlsx.py  ->  data/nj_campaign_money_export.xlsx
"""
import csv, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

D = json.load(open("data/site_data.json"))
ENT = {e["id"]: e for e in D["entities"]}
KEY = D["key"]
PARTY = {"D": "Democratic", "R": "Republican"}

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HEAD_FILL = PatternFill("solid", fgColor="E6EDFB")
SECTION = Font(bold=True, size=12, color="1D4FB8")
MONEY = '"$"#,##0.00'


def tab_name(county, party, q):
    return f"{county} {party} Q{q}"


def widths(ws, cols):
    for i, w in cols.items():
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()
summary = wb.active
summary.title = "Summary"

S_HEAD = ["County", "Party", "Committee", "Quarter", "Filed", "Amendment", "ELEC filing #", "Opening balance",
          "Total receipts", "Unitemized ($200 or less)", "Itemized contributions", "# contributions",
          "Total expenditures (filed)", "# expenditures", "Closing balance (cash on hand)", "Tab"]
summary.append(["NJ Campaign Money Map: county party committees, 2026 Q1–Q2 (ELEC Form R-3)"])
summary["A1"].font = TITLE
summary.append(["One tab per committee per quarter. Individuals' home addresses are omitted. Source PDFs and parser: "
                "github.com/SamSobolov/nj-campaign-money-map"])
summary.append([])
summary.append(S_HEAD)
for c in summary[4]:
    c.font = BOLD; c.fill = HEAD_FILL

for ck, com in sorted(D["committees"].items()):
    county, party = ck.split("|")
    for f in sorted(com["filings"], key=lambda f: f["quarter"]):
        q = f["quarter"]
        name = tab_name(county, party, q)
        ents = {e["id"] for e in D["entities"] if e["county"] == county and e["party"] == party}
        contribs = sorted((c for c in D["contributions"] if c["entityId"] in ents and c["q"] == q),
                          key=lambda c: (c["date"], c["name"]))
        expends = sorted((x for x in D["expenditures"] if x["entityId"] in ents and x["q"] == q),
                         key=lambda x: (x["date"], x["vendor"]))

        ws = wb.create_sheet(name)
        ws["A1"] = f'{com["name"]} · {county} County {PARTY[party]} · 2026 Q{q}'
        ws["A1"].font = TITLE
        ws["A2"] = (f'ELEC Form R-3 filing #{f["elecId"]} · filed {f["filed"]}'
                    + (f' · amendment #{f["amendment"]}' if f["amendment"] else ""))
        labels = ["Opening balance", "Total receipts", "Unitemized ($200 or less)", "Total expenditures",
                  "Closing balance (cash on hand)"]
        vals = [f["opening"], f["receipts"], f["unitemized"], f["expenditures"], f["closing"]]
        for i, (l, v) in enumerate(zip(labels, vals)):
            ws.cell(3, 1 + i * 2, l).font = BOLD
            c = ws.cell(4, 1 + i * 2, v); c.number_format = MONEY

        ws["A6"] = "CONTRIBUTIONS"; ws["A6"].font = SECTION
        ws["G6"] = "Itemized total"; ws["G6"].font = BOLD
        ws["H6"] = sum(c["amount"] for c in contribs); ws["H6"].number_format = MONEY
        ws["J6"] = "EXPENDITURES"; ws["J6"].font = SECTION
        ws["P6"] = "Total"; ws["P6"].font = BOLD
        ws["Q6"] = sum(x["amount"] for x in expends); ws["Q6"].number_format = MONEY

        ch = ["Date", "Contributor", "Type", "Employer", "Occupation", "Account", "In-kind", "Amount"]
        eh = ["Date", "Vendor", "Vendor address", "Purpose (as filed)", "Bucket", "ELEC schedule", "Account", "Amount"]
        for i, h in enumerate(ch):
            c = ws.cell(7, 1 + i, h); c.font = BOLD; c.fill = HEAD_FILL
        for i, h in enumerate(eh):
            c = ws.cell(7, 10 + i, h); c.font = BOLD; c.fill = HEAD_FILL
        for r, c in enumerate(contribs, start=8):
            row = [c["date"], c["name"], c["type"], c["employer"], c["occupation"], ENT[c["entityId"]]["account"],
                   "Yes" if c["inKind"] else "", c["amount"]]
            for i, v in enumerate(row):
                ws.cell(r, 1 + i, v)
            ws.cell(r, 8).number_format = MONEY
        for r, x in enumerate(expends, start=8):
            row = [x["date"], x["vendor"], x["address"], x["purpose"], KEY[x["code"]], x["schedule"],
                   ENT[x["entityId"]]["account"], x["amount"]]
            for i, v in enumerate(row):
                ws.cell(r, 10 + i, v)
            ws.cell(r, 17).number_format = MONEY
        if not contribs:
            ws["A8"] = "No itemized contributions this quarter."
        if not expends:
            ws["J8"] = "No itemized expenditures this quarter."
        ws.freeze_panes = "A8"
        widths(ws, {1: 12, 2: 34, 3: 18, 4: 30, 5: 20, 6: 28, 7: 8, 8: 13, 9: 3,
                    10: 12, 11: 32, 12: 36, 13: 40, 14: 24, 15: 9, 16: 28, 17: 13})

        summary.append([county, PARTY[party], com["name"], f"Q{q}", f["filed"], f["amendment"] or "", f["elecId"],
                        f["opening"], f["receipts"], f["unitemized"], sum(c["amount"] for c in contribs),
                        len(contribs), f["expenditures"], len(expends), f["closing"], name])
        row = summary.max_row
        link = summary.cell(row, 16)
        link.hyperlink = f"#'{name}'!A1"
        link.font = Font(color="1D4FB8", underline="single")
        for col in (8, 9, 10, 11, 13, 15):
            summary.cell(row, col).number_format = MONEY

summary.freeze_panes = "A5"
summary.auto_filter.ref = f"A4:P{summary.max_row}"
widths(summary, {1: 12, 2: 12, 3: 44, 4: 8, 5: 22, 6: 11, 7: 13, 8: 15, 9: 15, 10: 15, 11: 16, 12: 9, 13: 17,
                 14: 9, 15: 18, 16: 18})

key = wb.create_sheet("Bucket Key")
key.append(["Spending bucket", "ELEC purpose as filed", "Total", "# expenditures"])
for c in key[1]:
    c.font = BOLD; c.fill = HEAD_FILL
for row in csv.DictReader(open("data/bucket_mapping.csv")):
    key.append([row["bucket"], row["purpose_as_filed"], float(row["total"]), int(row["count"])])
    key.cell(key.max_row, 3).number_format = MONEY
key.freeze_panes = "A2"
widths(key, {1: 26, 2: 60, 3: 14, 4: 14})

notes = wb.create_sheet("Notes")
for line in [
    "Source: NJ ELEC Form R-3 quarterly reports for county party committees, 2026 Q1–Q2.",
    "Contributions = Schedule 1 (monetary, over $200) and Schedule 2 (in-kind). Unitemized ($200 or less) appears only as a total.",
    "Expenditures = Schedules 8 (operating), 10 (contributions to candidates/committees) and 11 (on behalf of candidates). Sub-payee breakdowns are not repeated.",
    "Voided checks appear as negative amounts, as filed. When a quarter was amended, the latest amendment is used.",
    "Contributor Type is inferred from employer/occupation and the name; ELEC's R-3 does not label it.",
    "Spending buckets are provisional, assigned from ELEC purpose text (see Bucket Key).",
    "Street addresses for individual donors and individual payees are omitted.",
]:
    notes.append([line])
notes.column_dimensions["A"].width = 140

wb.save("data/nj_campaign_money_export.xlsx")
print(len(wb.sheetnames), "tabs ->", "data/nj_campaign_money_export.xlsx")
